"""Streamlit UI for Malaria SSR1.6 validation.
Run: streamlit run malaria_streamlit_app.py
Keep Malaria_SSR1.6_v1(1).py in the same folder.
"""

from __future__ import annotations

import importlib.util
from io import BytesIO
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

BASE_DIR = Path(__file__).resolve().parent
RULE_FILE = BASE_DIR / "Malaria_SSR1.6_v1(1).py"
TEMPLATE_FILE = BASE_DIR / "Malaria_SSR1_6_Input_Template.xlsx"
LOCAL_PHOTO = BASE_DIR / "anopheles_mosquito.jpg"
PHOTO_URL = "https://commons.wikimedia.org/wiki/Special:FilePath/Anopheles_minimus.jpg?width=960"

st.set_page_config(page_title="Malaria SSR1.6 | Data Validation", page_icon="🦟", layout="wide")
st.markdown("""
<style>
.block-container {padding-top: 1.5rem;}
.hero-title {font-size:2rem; font-weight:750; color:#146451; margin-bottom:.3rem;}
.hero-sub {color:#586a68; font-size:1.0rem;}
</style>
""", unsafe_allow_html=True)


@st.cache_resource(show_spinner=False)
def load_rules(rule_path: str, modified_ns: int):
    spec = importlib.util.spec_from_file_location("malaria_ssr16_rules", rule_path)
    if spec is None or spec.loader is None:
        raise RuntimeError("Could not load malaria rules module.")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def read_sheet(upload, sheet_name):
    # Keep data as supplied, without allowing inferred numeric conversion of IDs.
    # Native Excel dates remain as timestamps; the rules decide acceptable date formats.
    return pd.read_excel(upload, sheet_name=sheet_name, dtype=object, engine="openpyxl")


def select_sheet(upload, key):
    if upload is None:
        return None
    xl = pd.ExcelFile(upload, engine="openpyxl")
    return st.selectbox("Worksheet", xl.sheet_names, key=key)


def to_excel(positive=None, aggregate=None, recheck=None):
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        if positive is not None:
            positive.to_excel(writer, sheet_name="Positive_Checked", index=False)
        if aggregate is not None:
            aggregate.to_excel(writer, sheet_name="Aggregate_Checked", index=False)
        if recheck is not None:
            rows = recheck.get("custom_rows")
            if rows:
                wb = writer.book
                ws = wb.create_sheet(recheck.get("sheet_title", "REACH_RECHECK")[:31])
                for row in rows:
                    ws.append([None if pd.isna(v) and not isinstance(v, str) else v for v in row])
                for cell in ws[1]:
                    cell.fill = PatternFill("solid", fgColor="146451")
                    cell.font = Font(color="FFFFFF", bold=True)
                for cell in ws[2]:
                    cell.fill = PatternFill("solid", fgColor="DCECE8")
                    cell.font = Font(bold=True)
                    cell.alignment = Alignment(wrap_text=True, vertical="center")
                ws.freeze_panes = "G3"
            else:
                recheck["file"].to_excel(writer, sheet_name="REACH_RECHECK", index=False)
        for ws in writer.book.worksheets:
            if ws.title != "REACH_RECHECK":
                ws.freeze_panes = "A2"
                ws.auto_filter.ref = ws.dimensions
                for cell in ws[1]:
                    cell.fill = PatternFill("solid", fgColor="146451")
                    cell.font = Font(color="FFFFFF", bold=True)
                for column in ws.columns:
                    letter = column[0].column_letter
                    ws.column_dimensions[letter].width = min(35, max(13, len(str(column[0].value or "")) + 3))
    return output.getvalue()


def count_comments(df):
    if "COMMENT" not in df:
        return 0
    comments = df["COMMENT"].fillna("").astype(str).str.strip()
    return int(comments.ne("").sum())


heading, photo = st.columns([3.2, 1], vertical_alignment="center")
with heading:
    st.markdown('<div class="hero-title">Malaria SSR1.6 — Data Validation</div>', unsafe_allow_html=True)
    st.markdown('<div class="hero-sub">Upload positive and aggregate malaria records, check data quality, calculate indicators, and export results.</div>', unsafe_allow_html=True)
with photo:
    # A local copy works offline; otherwise Streamlit loads the Wikimedia Commons photo.
    try:
        st.image(str(LOCAL_PHOTO) if LOCAL_PHOTO.exists() else PHOTO_URL,
                 caption="Anopheles minimus • Photo: CDC / James Gathany (public domain)",
                 use_container_width=True)
    except Exception:
        st.caption("Mosquito photo not available offline. Add anopheles_mosquito.jpg beside this app.")

st.divider()
if TEMPLATE_FILE.is_file():
    st.download_button("Download blank Excel input template", data=TEMPLATE_FILE.read_bytes(),
                       file_name=TEMPLATE_FILE.name,
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

if not RULE_FILE.is_file():
    st.error(f"Rule script missing. Place '{RULE_FILE.name}' beside this Streamlit app.")
    st.stop()

mode = st.radio("Run checks", ["Positive + Aggregate", "Positive only", "Aggregate only"],
                horizontal=True, index=0)
need_pos = mode != "Aggregate only"
need_agg = mode != "Positive only"

left, right = st.columns(2)
with left:
    st.subheader("Positive cases")
    pos_upload = st.file_uploader("Upload positive records (.xlsx)", type=["xlsx"],
                                  key="positive", disabled=not need_pos)
    pos_sheet = select_sheet(pos_upload, "positive_sheet") if need_pos else None
with right:
    st.subheader("Aggregate / total records")
    agg_upload = st.file_uploader("Upload aggregate records (.xlsx)", type=["xlsx"],
                                  key="aggregate", disabled=not need_agg)
    agg_sheet = select_sheet(agg_upload, "aggregate_sheet") if need_agg else None

can_run = (not need_pos or pos_upload is not None) and (not need_agg or agg_upload is not None)
run = st.button("Run malaria validation", type="primary", disabled=not can_run,
                use_container_width=True)

if run:
    try:
        rules = load_rules(str(RULE_FILE), RULE_FILE.stat().st_mtime_ns)
        positive_source = read_sheet(pos_upload, pos_sheet) if need_pos else None
        aggregate_source = read_sheet(agg_upload, agg_sheet) if need_agg else None
        with st.spinner("Checking records against SSR1.6 rules..."):
            checked_pos = rules.malaria_positive(positive_source) if need_pos else None
            checked_agg = rules.malaria_aggregate(aggregate_source) if need_agg else None
            recheck = None
            if checked_pos is not None and checked_agg is not None:
                recheck = rules.malaria_recheck(
                    checked_pos, checked_agg,
                    original_input_columns=set(positive_source.columns),
                )
            result_file = to_excel(checked_pos, checked_agg, recheck)
        st.session_state["result"] = (checked_pos, checked_agg, recheck, result_file)
    except Exception as exc:
        st.session_state.pop("result", None)
        st.error(f"Validation failed: {exc}")

if "result" in st.session_state:
    checked_pos, checked_agg, recheck, result_file = st.session_state["result"]
    st.success("Validation finished. Review findings and download the checked workbook.")
    metrics = st.columns(4)
    metrics[0].metric("Positive records", len(checked_pos) if checked_pos is not None else "—")
    metrics[1].metric("Positive rows with comments", count_comments(checked_pos) if checked_pos is not None else "—")
    metrics[2].metric("Aggregate records", len(checked_agg) if checked_agg is not None else "—")
    metrics[3].metric("Aggregate rows with comments", count_comments(checked_agg) if checked_agg is not None else "—")

    st.download_button("Download all validation results (.xlsx)", data=result_file,
                       file_name="Malaria_SSR1_6_Validation_Results.xlsx",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                       type="primary")

    tabs = []
    if checked_pos is not None:
        tabs.append(("Positive results", checked_pos))
    if checked_agg is not None:
        tabs.append(("Aggregate results", checked_agg))
    if recheck is not None:
        tabs.append(("REACH_RECHECK preview", recheck.get("display", pd.DataFrame())))
    tab_widgets = st.tabs([x[0] for x in tabs])
    for tab, (_, df) in zip(tab_widgets, tabs):
        with tab:
            if "COMMENT" in df.columns:
                only_findings = st.checkbox("Show rows with comments only", key=f"filter_{tab_widgets.index(tab)}")
                if only_findings:
                    df = df[df["COMMENT"].fillna("").astype(str).str.strip().ne("")]
            st.dataframe(df, use_container_width=True, hide_index=True)
    if recheck is not None:
        st.info("REACH_RECHECK includes Excel formulas. Open the downloaded workbook in Excel or LibreOffice to recalculate the comparison fields.")
